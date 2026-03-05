from langchain_core.tools import tool
import os

@tool
def modify_word_docx(file_path: str, font_changes: dict = None, add_content: str = None):
    """
    修改Word文档，包括字体样式调整和内容添加
    
    Args:
        file_path: Word文档路径
        font_changes: 字体修改设置，如{'font_name': '微软雅黑', 'font_size': 12}
        add_content: 要添加的内容
    
    Returns:
        dict: 操作结果
    """
    try:
        try:
            from docx import Document
            from docx.shared import Pt
            from docx.oxml.ns import qn
        except Exception:
            return {
                "success": False,
                "error": "缺少依赖 python-docx，请先安装后再使用该工具",
                "solution": "pip install python-docx"
            }
        # 检查文件是否存在
        if not os.path.exists(file_path):
            return {"success": False, "error": f"文件不存在: {file_path}"}
        
        # 读取文档
        doc = Document(file_path)
        
        # 修改字体样式
        if font_changes:
            font_name = font_changes.get('font_name')
            font_size = font_changes.get('font_size')
            
            # 遍历所有段落修改字体
            for paragraph in doc.paragraphs:
                for run in paragraph.runs:
                    if font_name:
                        run.font.name = font_name
                        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
                    if font_size:
                        run.font.size = Pt(font_size)
            
            # 遍历所有表格修改字体
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                if font_name:
                                    run.font.name = font_name
                                    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
                                if font_size:
                                    run.font.size = Pt(font_size)
        
        # 添加内容
        if add_content:
            doc.add_paragraph(add_content)
        
        # 保存文档
        root, ext = os.path.splitext(file_path)
        backup_path = f"{root}_backup{ext or '.docx'}"
        original_doc = Document(file_path)  # 重新加载原始文档以保存备份
        original_doc.save(backup_path)  # 保存备份
        
        doc.save(file_path)  # 保存修改后的文档
        
        return {
            "success": True, 
            "message": f"文档修改成功: {file_path}",
            "backup_created": backup_path
        }
    except Exception as e:
        return {"success": False, "error": str(e)}

if __name__ == "__main__":
    # 测试函数
    result = modify_word_docx("test.docx", {"font_name": "微软雅黑", "font_size": 12}, "测试添加内容")
    print(result)
