from langchain_core.tools import tool
import os
import tempfile
import urllib.request
from typing import Any, Dict, List, Optional
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor


@tool
def create_ppt(
    title: str,
    content: List[Any],
    output_path: str,
    theme: str = "professional"
) -> Dict[str, Any]:
    """
    创建生动形象的PPT文件，支持主题、图文混排与重点强调

    Args:
        title: PPT标题
        content: 幻灯片列表，元素可为字符串或字典
        output_path: 输出路径，需以 .pptx 结尾
        theme: 主题风格，professional | vivid | dark | fresh

    Returns:
        操作结果字典
    """
    if not output_path or not output_path.lower().endswith(".pptx"):
        return {"success": False, "error": "output_path 必须以 .pptx 结尾"}
    if content is None:
        content = []
    if not isinstance(content, list):
        return {"success": False, "error": "content 必须为列表"}
    if not title:
        title = "未命名演示文稿"

    theme_map = {
        "professional": {
            "title": RGBColor(40, 55, 71),
            "accent": RGBColor(52, 152, 219),
            "body": RGBColor(33, 33, 33),
            "background": RGBColor(255, 255, 255)
        },
        "vivid": {
            "title": RGBColor(41, 128, 185),
            "accent": RGBColor(231, 76, 60),
            "body": RGBColor(44, 62, 80),
            "background": RGBColor(255, 255, 255)
        },
        "dark": {
            "title": RGBColor(236, 240, 241),
            "accent": RGBColor(52, 152, 219),
            "body": RGBColor(236, 240, 241),
            "background": RGBColor(30, 30, 30)
        },
        "fresh": {
            "title": RGBColor(46, 204, 113),
            "accent": RGBColor(155, 89, 182),
            "body": RGBColor(44, 62, 80),
            "background": RGBColor(245, 251, 248)
        }
    }
    palette = theme_map.get(theme, theme_map["professional"])

    prs = Presentation()
    temp_files: List[str] = []
    errors: List[str] = []

    def apply_background(slide):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = palette["background"]

    def add_title(slide, text: str, is_cover: bool):
        if slide.shapes.title is not None:
            title_shape = slide.shapes.title
            title_shape.text = text
        else:
            title_shape = slide.shapes.add_textbox(int(prs.slide_width * 0.1), int(prs.slide_height * 0.08), int(prs.slide_width * 0.8), int(prs.slide_height * 0.15))
            title_shape.text_frame.text = text
        for paragraph in title_shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(48 if is_cover else 32)
                run.font.bold = True
                run.font.color.rgb = palette["title"]

    def add_body_text(slide, lines: List[str], left: int, top: int, width: int, height: int):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.clear()
        for idx, line in enumerate(lines):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = line
            p.font.size = Pt(20)
            p.font.color.rgb = palette["body"]

    def resolve_image(image_ref: str) -> Optional[str]:
        if not image_ref:
            return None
        if image_ref.lower().startswith("http://") or image_ref.lower().startswith("https://"):
            try:
                suffix = os.path.splitext(image_ref.split("?")[0])[1] or ".png"
                fd, path = tempfile.mkstemp(suffix=suffix)
                os.close(fd)
                urllib.request.urlretrieve(image_ref, path)
                temp_files.append(path)
                return path
            except Exception as e:
                errors.append(f"图片下载失败: {image_ref} - {e}")
                return None
        if not os.path.exists(image_ref):
            errors.append(f"图片不存在: {image_ref}")
            return None
        return image_ref

    cover = prs.slides.add_slide(prs.slide_layouts[0])
    apply_background(cover)
    add_title(cover, title, True)
    if cover.placeholders and len(cover.placeholders) > 1:
        subtitle_placeholder = cover.placeholders[1]
        subtitle_placeholder.text = " "
        for paragraph in subtitle_placeholder.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(18)
                run.font.color.rgb = palette["accent"]

    def parse_text_block(text: str):
        title_text = ""
        paragraphs_list: List[str] = []
        bullets_list: List[str] = []
        if not text:
            return title_text, paragraphs_list, bullets_list
        lines = [ln.strip() for ln in str(text).splitlines() if str(ln).strip()]
        if not lines:
            return title_text, paragraphs_list, bullets_list
        if len(lines) == 1:
            title_text = lines[0]
            return title_text, paragraphs_list, bullets_list
        title_text = lines[0]
        for line in lines[1:]:
            if line.startswith(("-", "•", "*")):
                bullets_list.append(line.lstrip("-•* ").strip())
            elif line[:2].isdigit() and line[2:3] in [".", "、"]:
                bullets_list.append(line[3:].strip())
            else:
                paragraphs_list.append(line)
        return title_text, paragraphs_list, bullets_list

    def normalize_list(value):
        if value is None:
            return []
        if isinstance(value, list):
            return [v for v in value if v is not None]
        return [value]

    def normalize_item(item_value):
        if isinstance(item_value, dict):
            title_value = item_value.get("title")
            paragraphs_value = item_value.get("paragraphs")
            bullets_value = item_value.get("bullets")
            highlights_value = item_value.get("highlights")
            images_value = item_value.get("images")
            if images_value is None:
                images_value = item_value.get("image") or item_value.get("image_url") or item_value.get("image_path") or item_value.get("img")
            if paragraphs_value is None:
                paragraphs_value = item_value.get("content") or item_value.get("text") or item_value.get("body") or item_value.get("desc") or item_value.get("description")
            if bullets_value is None:
                bullets_value = item_value.get("points") or item_value.get("items") or item_value.get("outline") or item_value.get("list")
            if highlights_value is None:
                highlights_value = item_value.get("key_points") or item_value.get("keypoints") or item_value.get("conclusion")
            if not title_value and isinstance(paragraphs_value, str):
                parsed_title, parsed_paragraphs, parsed_bullets = parse_text_block(paragraphs_value)
                if parsed_title:
                    title_value = parsed_title
                if parsed_paragraphs and paragraphs_value == item_value.get("paragraphs"):
                    paragraphs_value = parsed_paragraphs
                if parsed_bullets and bullets_value is None:
                    bullets_value = parsed_bullets
            return {
                "title": title_value,
                "paragraphs": normalize_list(paragraphs_value),
                "bullets": normalize_list(bullets_value),
                "highlights": normalize_list(highlights_value),
                "images": normalize_list(images_value)
            }
        if isinstance(item_value, str):
            title_value, paragraphs_value, bullets_value = parse_text_block(item_value)
            return {
                "title": title_value or item_value,
                "paragraphs": paragraphs_value,
                "bullets": bullets_value,
                "highlights": [],
                "images": []
            }
        return {
            "title": "内容",
            "paragraphs": [str(item_value)],
            "bullets": [],
            "highlights": [],
            "images": []
        }

    for item in content:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        apply_background(slide)
        normalized = normalize_item(item)
        slide_title = normalized.get("title") or " "
        add_title(slide, slide_title, False)

        bullets = normalized.get("bullets") or []
        paragraphs = normalized.get("paragraphs") or []
        highlights = normalized.get("highlights") or []
        images = normalized.get("images") or []

        text_lines: List[str] = []
        for text in paragraphs:
            if text:
                text_lines.append(str(text))
        for bullet in bullets:
            if bullet:
                text_lines.append(f"• {bullet}")
        for highlight in highlights:
            if highlight:
                text_lines.append(f"★ {highlight}")

        has_images = len(images) > 0
        if text_lines:
            text_left = int(prs.slide_width * 0.08)
            text_top = int(prs.slide_height * 0.22)
            text_width = int(prs.slide_width * (0.52 if has_images else 0.84))
            text_height = int(prs.slide_height * 0.65)
            add_body_text(slide, text_lines, text_left, text_top, text_width, text_height)

        if has_images:
            for idx, img_ref in enumerate(images[:2]):
                img_path = resolve_image(str(img_ref))
                if not img_path:
                    continue
                left = int(prs.slide_width * (0.6 if idx == 0 else 0.12))
                top = int(prs.slide_height * (0.22 if idx == 0 else 0.62))
                width = int(prs.slide_width * (0.35 if idx == 0 else 0.3))
                try:
                    slide.shapes.add_picture(img_path, left, top, width=width)
                except Exception as e:
                    errors.append(f"图片插入失败: {img_ref} - {e}")

    output_dir = os.path.dirname(os.path.abspath(output_path))
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    try:
        prs.save(output_path)
    except Exception as e:
        return {"success": False, "error": f"PPT保存失败: {e}"}
    finally:
        for path in temp_files:
            try:
                os.remove(path)
            except Exception:
                pass

    return {
        "success": True,
        "output_path": output_path,
        "slides_count": len(prs.slides),
        "errors": errors
    }
